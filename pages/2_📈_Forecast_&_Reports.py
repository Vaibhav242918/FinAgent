import streamlit as st
import pandas as pd
import numpy as np
import sqlite3
import plotly.express as px
import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from datetime import datetime, timedelta

st.set_page_config(page_title="FinAgent - Forecast", page_icon="📈", layout="wide")
st.title("📈 Machine Learning Forecast & Reports")
st.markdown("Predict future spending trends and export enterprise-grade reports.")
st.divider()

# --- NEW: DASHBOARD CONTROLS (CLEAR DATA) ---
with st.sidebar:
    st.header("⚙️ Dashboard Controls")
    st.markdown("Use this to wipe your old database records and reset the charts.")
    if st.button("🗑️ Clear All Expense Data", use_container_width=True, type="primary"):
        try:
            conn = sqlite3.connect("finagent.db")
            cursor = conn.cursor()
            cursor.execute("DELETE FROM expenses")  # Wipes the table completely
            conn.commit()
            conn.close()
            st.success("✅ All database records cleared!")
            st.rerun()  # Instantly refreshes the page to remove the chart
        except Exception as e:
            st.error(f"Database error: {e}")

# 1. Fetch Data from SQLite Database
def get_historical_data():
    conn = sqlite3.connect("finagent.db")
    df = pd.read_sql_query("SELECT date, category, amount FROM expenses ORDER BY date", conn)
    conn.close()
    return df

df = get_historical_data()

if df.empty:
    st.warning("⚠️ No data available. Log some expenses or upload a file first!")
else:
    # --- OPTION 1: PURE MATH FORECAST ---
    st.header("🤖 30-Day Predictive Spending Model")
    
    # Data Preparation for ML
    df['date'] = pd.to_datetime(df['date'])
    df_daily = df.groupby('date').sum(numeric_only=True).reset_index()
    
    if len(df_daily) < 3:
        st.info("Not enough days of data to run the predictive model. Need at least 3 days of logged expenses.")
    else:
        # Convert dates to a numerical format (ordinal)
        df_daily['date_ordinal'] = df_daily['date'].map(pd.Timestamp.toordinal)
        
        x = df_daily['date_ordinal']
        y = df_daily['amount']
        
        x_mean = x.mean()
        y_mean = y.mean()
        
        # Calculate Slope (m) and Intercept (c)
        m = ((x - x_mean) * (y - y_mean)).sum() / ((x - x_mean)**2).sum()
        c = y_mean - m * x_mean
        
        last_date = df_daily['date'].max()
        future_dates = [last_date + timedelta(days=i) for i in range(1, 31)]
        
        future_ordinals = pd.Series([d.toordinal() for d in future_dates])
        predictions = m * future_ordinals + c
        
        future_df = pd.DataFrame({'date': future_dates, 'amount': predictions, 'Type': 'Predicted'})
        df_daily['Type'] = 'Historical'
        
        combined_df = pd.concat([df_daily[['date', 'amount', 'Type']], future_df])
        combined_df['amount'] = combined_df['amount'].apply(lambda val: max(0, val))

        # Clean Plotly Visualization
        fig = px.scatter(
            combined_df, 
            x="date", 
            y="amount", 
            color="Type", 
            title="Historical Spending vs. 30-Day Forecast"
        )
        fig.update_traces(marker=dict(size=8))
        st.plotly_chart(fig, use_container_width=True)

    st.divider()

    # --- OPTION 3: EXPORTABLE EXCEL REPORTS ---
    st.header("📑 Enterprise Report Generation")
    st.markdown("Generate a formatted `.xlsx` financial statement of your raw data.")

    def generate_excel_report(dataframe):
        output = io.BytesIO()
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Monthly Expense Report"

        ws.merge_cells('A1:C1')
        title_cell = ws['A1']
        title_cell.value = f"FinAgent - Financial Report ({datetime.now().strftime('%Y-%m-%d')})"
        title_cell.font = Font(size=14, bold=True, color="FFFFFF")
        title_cell.fill = PatternFill("solid", fgColor="1F4E79")
        title_cell.alignment = Alignment(horizontal="center")

        headers = ["Date", "Category", "Amount (INR)"]
        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=3, column=col_num, value=header)
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="D3D3D3")
        
        for row_num, row_data in enumerate(dataframe.values, 4):
            date_val = row_data[0]
            if isinstance(date_val, pd.Timestamp):
                date_val = date_val.strftime('%Y-%m-%d')
            ws.cell(row=row_num, column=1, value=date_val)
            ws.cell(row=row_num, column=2, value=row_data[1])
            amount_cell = ws.cell(row=row_num, column=3, value=row_data[2])
            amount_cell.number_format = '₹#,##0.00'

        ws.column_dimensions['A'].width = 15
        ws.column_dimensions['B'].width = 20
        ws.column_dimensions['C'].width = 15

        wb.save(output)
        return output.getvalue()

    excel_data = generate_excel_report(df)
    
    st.download_button(
        label="📥 Download Formatted Excel Report",
        data=excel_data,
        file_name=f"FinAgent_Report_{datetime.now().strftime('%Y%m%d')}.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        type="primary"
    )